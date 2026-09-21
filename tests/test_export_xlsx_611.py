# -*- coding: utf-8 -*-
"""6.1.2 (B1): 数据导出升级 — Excel 口径说明 sheet + 中文文件名"""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from api.v1.export import router
from auth import get_non_guest_user


@pytest.fixture()
def client():
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_non_guest_user] = lambda: {"username": "tester"}
    return TestClient(app)


def test_xlsx_has_chinese_filename(client):
    r = client.get("/api/data/export/excel")
    assert r.status_code == 200
    disp = r.headers.get("content-disposition", "")
    assert "UTF-8''" in disp and "%E9%87%8F%E5%8C%96" in disp and ".xlsx" in disp


def test_xlsx_contains_meta_sheet(client):
    r = client.get("/api/data/export/excel")
    assert r.status_code == 200
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "口径说明" in wb.sheetnames, f"缺少口径说明 sheet: {wb.sheetnames}"
    ws = wb["口径说明"]
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("导出时间" in v for v in flat)
    assert any("免责声明" in v for v in flat)
    assert any("数据来源" in v for v in flat)


def test_csv_has_chinese_filename(client):
    r = client.get("/api/data/export/csv?type=strategies")
    assert r.status_code == 200
    disp = r.headers.get("content-disposition", "")
    assert "UTF-8''" in disp and "%E9%87%8F%E5%8C%96" in disp and ".csv" in disp
